"""Geometry-aware DOCX layout checks for OCR-004F."""

from __future__ import annotations

from pathlib import Path

from docx import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL

from ocr_document_bot.document_model import BoundingBox, Document, OCRProvenance, Page, Point, Table, TableCell, TableRow
from ocr_document_bot.renderers import DocxRenderer


def _provenance(order: int) -> OCRProvenance:
    return OCRProvenance(provider="test", source_page_number=0, source_order=order, source_type="visual_table")


def _box(left: int, top: int, right: int, bottom: int) -> BoundingBox:
    return BoundingBox(vertices=(Point(x=left, y=top), Point(x=right, y=bottom)))


def test_wide_geometry_uses_landscape_and_preserves_column_ratios(tmp_path: Path) -> None:
    # Original coordinates encode [100, 300, 50, 350] column proportions.
    table = Table(
        id="wide", geometry=_box(0, 0, 800, 200), provenance=_provenance(0), row_count=2, column_count=4,
        rows=[
            TableRow(index=0, cells=[
                TableCell(text="Merged header", row_index=0, column_index=0, column_span=2, geometry=_box(0, 0, 400, 100), provenance=_provenance(0)),
                TableCell(text="Technical", row_index=0, column_index=2, geometry=_box(400, 0, 450, 100), provenance=_provenance(1)),
                TableCell(text="Notes", row_index=0, column_index=3, geometry=_box(450, 0, 800, 100), provenance=_provenance(2)),
            ]),
            TableRow(index=1, cells=[
                TableCell(text="A", row_index=1, column_index=0, geometry=_box(0, 100, 100, 200), provenance=_provenance(3)),
                TableCell(text="B", row_index=1, column_index=1, geometry=_box(100, 100, 400, 200), provenance=_provenance(4)),
                TableCell(text="1", row_index=1, column_index=2, geometry=_box(400, 100, 450, 200), provenance=_provenance(5)),
                TableCell(text="C", row_index=1, column_index=3, geometry=_box(450, 100, 800, 200), provenance=_provenance(6)),
            ]),
        ],
    )
    output = DocxRenderer().save(Document(source_provider="test", pages=[Page(number=0, elements=[table])]), tmp_path / "wide.docx")
    reopened = WordDocument(output)
    word_table = reopened.tables[0]
    widths = [column.width.twips for column in word_table.columns]

    assert reopened.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert len(word_table.rows) == 2 and len(word_table.columns) == 4
    assert widths[2] < widths[0] < widths[1] < widths[3]
    assert word_table.cell(0, 0)._tc is word_table.cell(0, 1)._tc
    assert word_table.cell(0, 0).vertical_alignment == WD_ALIGN_VERTICAL.TOP
    assert word_table.cell(0, 0).paragraphs[0].runs[0].font.size.pt == 8.5
