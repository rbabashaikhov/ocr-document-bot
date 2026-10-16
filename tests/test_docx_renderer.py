import json
from pathlib import Path

import pytest
from docx import Document as WordDocument

from ocr_document_bot.document_model import Document, Heading, OCRProvenance, Page, Paragraph, Table, TableCell, TableRow
from ocr_document_bot.providers.yandex_document_mapper import document_from_yandex_responses
from ocr_document_bot.renderers import DocxRenderer


def origin(order: int, source_type: str = "test") -> OCRProvenance:
    return OCRProvenance(provider="test", source_page_number=0, source_order=order, source_type=source_type)


def sample_table() -> Table:
    return Table(
        id="table", provenance=origin(1, "table"), row_count=2, column_count=2,
        rows=[
            TableRow(index=0, cells=[TableCell(text="Merged", row_index=0, column_index=0, column_span=2, provenance=origin(0, "cell"))]),
            TableRow(index=1, cells=[TableCell(text="", row_index=1, column_index=0, provenance=origin(1, "cell")), TableCell(text="Value\nwith break", row_index=1, column_index=1, provenance=origin(2, "cell"))]),
        ],
    )


def test_renders_paragraphs_headings_tables_and_reopens_docx(tmp_path: Path) -> None:
    document = Document(
        source_provider="test",
        pages=[Page(number=0, elements=[
            Heading(id="heading", text="Заголовок", level=1, provenance=origin(0)),
            Paragraph(id="before", text="Первый абзац\nс переносом", provenance=origin(1)),
            sample_table(),
            Paragraph(id="after", text="После таблицы", provenance=origin(3)),
        ])],
    )
    output = DocxRenderer().save(document, tmp_path / "result.docx")
    reopened = WordDocument(output)
    assert reopened.paragraphs[0].text == "Заголовок"
    assert reopened.paragraphs[0].style.name.startswith("Heading")
    assert reopened.paragraphs[1].text == "Первый абзац\nс переносом"
    assert reopened.paragraphs[-1].text == "После таблицы"
    assert len(reopened.tables) == 1
    table = reopened.tables[0]
    assert len(table.rows) == 2 and len(table.columns) == 2
    assert table.cell(0, 0).text == "Merged"
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(1, 0).text == ""
    assert table.cell(1, 1).text == "Value\nwith break"


def test_renders_page_breaks_between_document_model_pages(tmp_path: Path) -> None:
    document = Document(source_provider="test", pages=[
        Page(number=0, elements=[Paragraph(id="first", text="one", provenance=origin(0))]),
        Page(number=1, elements=[Paragraph(id="second", text="two", provenance=origin(0))]),
    ])
    reopened = WordDocument(DocxRenderer().save(document, tmp_path / "pages.docx"))
    assert [paragraph.text for paragraph in reopened.paragraphs if paragraph.text] == ["one", "two"]
    assert "w:type=\"page\"" in reopened._element.xml


def test_rejects_invalid_overlapping_cell_spans() -> None:
    table = Table(id="bad", provenance=origin(0), row_count=1, column_count=2, rows=[TableRow(index=0, cells=[
        TableCell(text="a", row_index=0, column_index=0, column_span=2, provenance=origin(0)),
        TableCell(text="b", row_index=0, column_index=1, provenance=origin(1)),
    ])])
    with pytest.raises(ValueError, match="overlaps"):
        DocxRenderer().render(Document(source_provider="test", pages=[Page(number=0, elements=[table])]))


def test_renders_real_native_6x2_fixture_without_duplicate_table_paragraphs(tmp_path: Path) -> None:
    fixture = Path("fixtures/vision/ocr-002b-mixed.json")
    if not fixture.exists():
        pytest.skip("local OCR-002B fixture unavailable")
    model = document_from_yandex_responses([json.loads(fixture.read_text())], model="markdown")
    reopened = WordDocument(DocxRenderer().save(model, tmp_path / "mixed.docx"))
    table = reopened.tables[0]
    assert len(table.rows) == 6 and len(table.columns) == 2
    cell_texts = [cell.text for row in table.rows for cell in row.cells]
    assert not any(paragraph.text in cell_texts for paragraph in reopened.paragraphs if paragraph.text)
