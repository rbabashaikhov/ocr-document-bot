"""Opt-in integration tests for visual ruled-table recovery."""

from __future__ import annotations

from pathlib import Path
import json
from collections import Counter

import cv2
import numpy as np
from docx import Document as WordDocument

from ocr_document_bot.document_model import Paragraph, Table, VisualTableDetector
from ocr_document_bot.providers.yandex_document_mapper import document_from_yandex_responses
from ocr_document_bot.renderers import DocxRenderer


def _image(tmp_path: Path, *, broken: bool = False, skew: float = 0) -> Path:
    image = np.full((180, 300, 3), 255, dtype=np.uint8)
    for x in (30, 110, 190, 270):
        cv2.line(image, (x, 30), (x, 150), (0, 0, 0), 2)
    for y in (30, 90, 150):
        cv2.line(image, (30, y), (270, y), (0, 0, 0), 2)
    # Top-left cell spans two columns.
    cv2.line(image, (110, 30), (110, 90), (255, 255, 255), 4)
    if broken:
        cv2.line(image, (270, 50), (270, 130), (255, 255, 255), 4)
    if skew:
        matrix = cv2.getRotationMatrix2D((150, 90), skew, 1.0)
        image = cv2.warpAffine(image, matrix, (300, 180), borderValue=(255, 255, 255))
    path = tmp_path / "table.png"
    assert cv2.imwrite(str(path), image)
    return path


def _block(x: int, y: int, text: str) -> dict[str, object]:
    return {"layoutType": "LAYOUT_TYPE_TEXT", "boundingBox": {"vertices": [{"x": x - 5, "y": y - 5}, {"x": x + 5, "y": y + 5}]}, "lines": [{"text": text}]}


def _payload() -> dict[str, object]:
    return {"textAnnotation": {"width": "300", "height": "180", "blocks": [
        _block(10, 10, "before"), _block(150, 60, "merged"), _block(70, 120, "left"),
        _block(150, 120, "middle"), _block(230, 120, "right"), _block(10, 170, "after"),
    ], "tables": []}}


def test_flag_off_keeps_existing_nonvisual_behavior(tmp_path: Path) -> None:
    baseline = document_from_yandex_responses([_payload()])
    model = document_from_yandex_responses([_payload()], source_images=[_image(tmp_path)], enable_visual_table_detection=False)

    assert not any(isinstance(element, Table) for element in model.pages[0].elements)
    assert model.model_dump() == baseline.model_dump()


def test_flag_on_creates_table_preserves_order_and_renders_docx(tmp_path: Path) -> None:
    model = document_from_yandex_responses([_payload()], source_images=[_image(tmp_path)], enable_visual_table_detection=True)
    elements = model.pages[0].elements
    table = next(element for element in elements if isinstance(element, Table))

    assert [element.kind for element in elements] == ["paragraph", "table", "paragraph"]
    assert (table.row_count, table.column_count) == (2, 3)
    assert table.rows[0].cells[0].column_span == 2
    assert [cell.text for row in table.rows for cell in row.cells] == ["merged", "", "left", "middle", "right"]
    reopened = WordDocument(DocxRenderer().save(model, tmp_path / "visual.docx"))
    assert len(reopened.tables) == 1
    assert reopened.tables[0].cell(0, 0)._tc is reopened.tables[0].cell(0, 1)._tc
    assert [paragraph.text for paragraph in reopened.paragraphs if paragraph.text] == ["before", "after"]


def test_native_table_has_priority_even_when_visual_flag_is_enabled(tmp_path: Path) -> None:
    payload = _payload()
    payload["textAnnotation"]["tables"] = [{"boundingBox": {"vertices": [{"x": "30", "y": "30"}, {"x": "270", "y": "150"}]}, "rowCount": "1", "columnCount": "1", "cells": [{"rowIndex": "0", "columnIndex": "0", "text": "native"}]}]

    model = document_from_yandex_responses([payload], source_images=[_image(tmp_path)], enable_visual_table_detection=True)
    tables = [element for element in model.pages[0].elements if isinstance(element, Table)]

    assert len(tables) == 1
    assert tables[0].provenance.source_type == "table"
    assert tables[0].rows[0].cells[0].text == "native"


def test_visual_candidate_keeps_all_source_text(tmp_path: Path) -> None:
    model = document_from_yandex_responses([_payload()], source_images=[_image(tmp_path, broken=True)], enable_visual_table_detection=True)

    assert not any(isinstance(element, Table) and element.provenance.source_type == "visual_table" for element in model.pages[0].elements)
    assert {element.text for element in model.pages[0].elements if isinstance(element, Paragraph)} == {"before", "merged", "left", "middle", "right", "after"}


def test_document_without_table_keeps_paragraphs(tmp_path: Path) -> None:
    image = np.full((180, 300, 3), 255, dtype=np.uint8)
    path = tmp_path / "plain.png"
    assert cv2.imwrite(str(path), image)
    model = document_from_yandex_responses([_payload()], source_images=[path], enable_visual_table_detection=True)

    assert not any(isinstance(element, Table) and element.provenance.source_type == "visual_table" for element in model.pages[0].elements)
    assert {element.text for element in model.pages[0].elements if isinstance(element, Paragraph)} == {"before", "merged", "left", "middle", "right", "after"}


def test_production_detector_handles_skew_and_multiple_tables(tmp_path: Path) -> None:
    path = _image(tmp_path, skew=3)
    detector = VisualTableDetector()
    skewed = detector.detect(path, [(150, 60), (70, 120), (150, 120), (230, 120)], ocr_width=300, ocr_height=180)

    assert skewed.status == "detected"
    assert abs(skewed.deskew_degrees) >= 1
    assert len(skewed.tables) == 1

    image = np.full((180, 620, 3), 255, dtype=np.uint8)
    for origin in (10, 350):
        for x in (origin, origin + 80, origin + 160, origin + 240):
            cv2.line(image, (x, 20), (x, 140), (0, 0, 0), 2)
        for y in (20, 80, 140):
            cv2.line(image, (origin, y), (origin + 240, y), (0, 0, 0), 2)
    second = tmp_path / "two.png"
    assert cv2.imwrite(str(second), image)
    two = detector.detect(second, [(50, 50), (390, 50)], ocr_width=620, ocr_height=180)
    assert two.status == "detected"
    assert len(two.tables) == 2


def test_local_complex_fixture_renders_before_after_without_duplicate_text(tmp_path: Path) -> None:
    fixture, image = Path("fixtures/vision/ocr-002b-table.json"), Path("tests/samples/table.png")
    if not fixture.exists() or not image.exists():
        return
    payload = json.loads(fixture.read_text(encoding="utf-8"))
    before = document_from_yandex_responses([payload], source_images=[image], enable_visual_table_detection=False)
    after = document_from_yandex_responses([payload], source_images=[image], enable_visual_table_detection=True)

    assert not any(isinstance(element, Table) for element in before.pages[0].elements)
    table = next(element for element in after.pages[0].elements if isinstance(element, Table))
    assert (table.row_count, table.column_count) == (5, 8)
    root = payload.get("result", payload)
    source_text = ["\n".join(line.get("text", "") for line in block.get("lines", [])) for block in root.get("textAnnotation", {}).get("blocks", [])]
    table_text = [cell.text for row in table.rows for cell in row.cells if cell.text]
    assert len(source_text) == 21
    assert all(any(text in rendered for rendered in table_text) for text in source_text)
    assert not any(isinstance(element, Paragraph) and element.text in source_text for element in after.pages[0].elements)
    reopened = WordDocument(DocxRenderer().save(after, tmp_path / "complex.docx"))
    assert len(reopened.tables) == 1
    seen_cells: set[str] = set()
    rendered_lines: list[str] = []
    for row in reopened.tables[0].rows:
        for cell in row.cells:
            # python-docx repeats a merged Word cell for every covered grid
            # position; its XML path identifies the one physical cell.
            path = cell._tc.getroottree().getpath(cell._tc)
            if path not in seen_cells:
                seen_cells.add(path)
                rendered_lines.extend(cell.text.splitlines())
    source_lines = [part for block in root.get("textAnnotation", {}).get("blocks", []) for line in block.get("lines", []) for part in str(line.get("text", "")).splitlines()]
    assert len(seen_cells) == 29
    assert Counter(rendered_lines) == Counter(source_lines)
