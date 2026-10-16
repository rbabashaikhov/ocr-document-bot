"""Regression coverage for the two OCR-006C mixed-PDF defects."""

from __future__ import annotations

from io import BytesIO

import fitz
from docx import Document as WordDocument
from PIL import Image

from ocr_document_bot.document_model import Paragraph, Table
from ocr_document_bot.domain import InputFile
from ocr_document_bot.ocr_engine import DocumentOCRService
from ocr_document_bot.preprocessing import InputPreprocessor
from ocr_document_bot.providers.yandex_document_mapper import document_from_yandex_responses
from ocr_document_bot.renderers import DocxRenderer


def _box(x1: int, y1: int, x2: int, y2: int) -> dict[str, object]:
    return {"vertices": [{"x": str(x1), "y": str(y1)}, {"x": str(x2), "y": str(y2)}]}


def _ocr_table_with_separate_numbers() -> dict[str, object]:
    values = (("Утренняя", "7", "4"), ("Дневная", "6", "4"), ("Вечерняя", "5", "4"))
    blocks = []
    cells = []
    for column, title in enumerate(("Shift", "Participants", "Hours")):
        cells.append({"rowIndex": "0", "columnIndex": str(column), "text": title, "boundingBox": _box(column * 100, 0, (column + 1) * 100, 25)})
    for row, values_row in enumerate(values, start=1):
        for column, value in enumerate(values_row):
            x1, y1 = column * 100, row * 25
            # Reproduces the observed provider shape: native cells omit numeric
            # strings that are nevertheless present as positioned OCR blocks.
            cells.append({"rowIndex": str(row), "columnIndex": str(column), "text": value if column == 0 else "", "boundingBox": _box(x1, y1, x1 + 100, y1 + 25)})
            blocks.append({"boundingBox": _box(x1 + 10, y1 + 5, x1 + 90, y1 + 20), "lines": [{"text": value}]})
    return {"textAnnotation": {"width": "300", "height": "100", "blocks": blocks, "tables": [{"rowCount": "4", "columnCount": "3", "boundingBox": _box(0, 0, 300, 100), "cells": cells}]}}


def _three_page_mixed_pdf() -> bytes:
    document = fitz.open()
    first = document.new_page()
    first.insert_text((72, 72), "First page text is intentionally long enough for native PDF extraction.")

    image = BytesIO()
    Image.new("RGB", (160, 90), "white").save(image, format="PNG")
    scanned = document.new_page()
    scanned.insert_image(scanned.rect, stream=image.getvalue())

    page = document.new_page()
    page.insert_text((72, 72), "Third page has a native table and enough text for extraction.")
    left, top, width, row_height = 72, 150, 360, 30
    for row in range(5):
        page.draw_line((left, top + row * row_height), (left + width, top + row * row_height))
    for column in range(4):
        page.draw_line((left + column * width / 3, top), (left + column * width / 3, top + 4 * row_height))
    for row, values in enumerate((("Metric", "Plan", "Actual"), ("Pages", "3", "3"), ("Tables", "2", "2"), ("Errors", "0", "0"))):
        for column, value in enumerate(values):
            page.insert_text((left + column * width / 3 + 8, top + (row + 1) * row_height - 10), value)
    page.insert_text((72, 330), "Final note remains outside the table and must stay after it.")
    result = document.tobytes()
    document.close()
    return result


class MockVision:
    def __init__(self) -> None:
        self.calls = 0

    def recognize(self, input_file: InputFile, *, language_codes: tuple[str, ...], model: str) -> dict[str, object]:
        self.calls += 1
        return _ocr_table_with_separate_numbers()


def _cell_values(table: Table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def test_separate_numeric_ocr_blocks_are_retained_in_native_table_and_docx(tmp_path) -> None:
    document = document_from_yandex_responses([_ocr_table_with_separate_numbers()])
    table = document.pages[0].elements[0]

    assert isinstance(table, Table)
    assert _cell_values(table) == [["Shift", "Participants", "Hours"], ["Утренняя", "7", "4"], ["Дневная", "6", "4"], ["Вечерняя", "5", "4"]]
    assert not [element for element in document.pages[0].elements if isinstance(element, Paragraph)]

    output = tmp_path / "numbers.docx"
    DocxRenderer().save(document, output)
    reopened = WordDocument(output)
    assert [[cell.text for cell in row.cells] for row in reopened.tables[0].rows] == _cell_values(table)


def test_three_page_mixed_pdf_preserves_two_editable_tables_and_order(tmp_path) -> None:
    vision = MockVision()
    service = DocumentOCRService(vision, preprocessor=InputPreprocessor())

    document = service.process(InputFile("mixed.pdf", "application/pdf", _three_page_mixed_pdf(), "telegram"))

    assert [page.number for page in document.pages] == [0, 1, 2]
    assert vision.calls == 1
    tables = [element for page in document.pages for element in page.elements if isinstance(element, Table)]
    assert [table.row_count for table in tables] == [4, 4]
    assert [table.column_count for table in tables] == [3, 3]
    assert _cell_values(tables[0])[1:] == [["Утренняя", "7", "4"], ["Дневная", "6", "4"], ["Вечерняя", "5", "4"]]
    assert _cell_values(tables[1]) == [["Metric", "Plan", "Actual"], ["Pages", "3", "3"], ["Tables", "2", "2"], ["Errors", "0", "0"]]
    page_three = document.pages[2].elements
    table_index = next(index for index, element in enumerate(page_three) if isinstance(element, Table))
    assert any(isinstance(element, Paragraph) and "Final note" in element.text for element in page_three[table_index + 1 :])

    output = tmp_path / "mixed.docx"
    DocxRenderer().save(document, output)
    reopened = WordDocument(output)
    assert len(reopened.tables) == 2
    assert [(len(table.rows), len(table.columns)) for table in reopened.tables] == [(4, 3), (4, 3)]
    assert [[cell.text for cell in row.cells] for row in reopened.tables[0].rows] == _cell_values(tables[0])
    assert [[cell.text for cell in row.cells] for row in reopened.tables[1].rows] == _cell_values(tables[1])
    assert sum(1 for paragraph in reopened.paragraphs for run in paragraph.runs if "w:type=\"page\"" in run._r.xml) == 2
